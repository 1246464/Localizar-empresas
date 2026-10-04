"""Localizador de empresas: consultas HTTP e normalização dos resultados."""
import math
import re
import requests

COLUNAS = ["Nome", "Endereço", "Avaliação", "Avaliações", "Latitude", "Longitude"]
SEGMENTOS = {
    "Restaurantes": ("restaurant", "amenity", "restaurant"),
    "Cafeterias": ("cafe", "amenity", "cafe"),
    "Farmácias": ("pharmacy", "amenity", "pharmacy"),
    "Supermercados": ("supermarket", "shop", "supermarket"),
    "Hotéis": ("hotel", "tourism", "hotel"),
    "Academias": ("gym", "leisure", "fitness_centre"),
    "Bancos": ("bank", "amenity", "bank"),
    "Hospitais": ("hospital", "amenity", "hospital"),
}


def coordenadas(latitude, longitude):
    lat = float(str(latitude).replace(",", "."))
    lon = float(str(longitude).replace(",", "."))
    if not (math.isfinite(lat) and math.isfinite(lon) and -90 <= lat <= 90 and -180 <= lon <= 180):
        raise ValueError("Informe latitude entre -90 e 90 e longitude entre -180 e 180.")
    return lat, lon


def consultar(method, url, **kwargs):
    """Não inclua URLs, tokens ou respostas do servidor nas mensagens de erro."""
    try:
        response = requests.request(method, url, timeout=(10, 40), **kwargs)
        if response.status_code in (401, 403):
            raise ValueError("Acesso negado. Confira a chave e as permissões do provedor.")
        if response.status_code == 429:
            raise ValueError("Limite de consultas atingido. Aguarde antes de tentar novamente.")
        response.raise_for_status()
        return response.json()
    except requests.Timeout:
        raise ValueError("O provedor demorou a responder. Tente um raio menor ou tente novamente.") from None
    except requests.exceptions.JSONDecodeError:
        raise ValueError("O provedor retornou uma resposta inválida.") from None
    except requests.RequestException:
        raise ValueError("Não foi possível consultar o provedor. Verifique a conexão e os parâmetros.") from None


def buscar(api, chave, raio, segmento, latitude, longitude):
    lat, lon = coordenadas(latitude, longitude)
    raio = int(raio)
    if not 1 <= raio <= (40000 if api == "Yelp" else 50000):
        raise ValueError("Raio fora do limite do provedor.")
    termo, osm_key, osm_value = SEGMENTOS.get(segmento, (segmento, "amenity", segmento))
    if api != "OSM" and not chave.strip():
        raise ValueError("Informe a chave de acesso do provedor.")
    rows = []
    if api == "Google":
        data = consultar("POST", "https://places.googleapis.com/v1/places:searchNearby",
            headers={"X-Goog-Api-Key": chave, "X-Goog-FieldMask": "places.displayName,places.formattedAddress,places.rating,places.userRatingCount,places.location"},
            json={"includedTypes": [termo], "maxResultCount": 20, "languageCode": "pt-BR",
                  "locationRestriction": {"circle": {"center": {"latitude": lat, "longitude": lon}, "radius": raio}}})
        for p in data.get("places", []):
            loc = p.get("location") or {}
            rows.append([p.get("displayName", {}).get("text", "Sem nome"), p.get("formattedAddress"), p.get("rating"), p.get("userRatingCount"), loc.get("latitude"), loc.get("longitude")])
    elif api == "Yelp":
        data = consultar("GET", "https://api.yelp.com/v3/businesses/search",
            headers={"Authorization": f"Bearer {chave}"},
            params={"term": termo, "latitude": lat, "longitude": lon, "radius": raio, "limit": 50})
        for p in data.get("businesses", []):
            loc = p.get("coordinates") or {}
            address = p.get("location") or {}
            rows.append([p.get("name", "Sem nome"), ", ".join(address.get("display_address") or []) or address.get("address1"), p.get("rating"), p.get("review_count"), loc.get("latitude"), loc.get("longitude")])
    elif api == "Foursquare":
        data = consultar("GET", "https://places-api.foursquare.com/places/search",
            headers={"Authorization": f"Bearer {chave}", "X-Places-Api-Version": "2025-06-17"},
            params={"ll": f"{lat},{lon}", "radius": raio, "query": termo, "limit": 50})
        for p in data.get("results", []):
            rows.append([p.get("name", "Sem nome"), (p.get("location") or {}).get("formatted_address"), None, None, p.get("latitude"), p.get("longitude")])
    elif api == "OSM":
        if not re.fullmatch(r"[a-z_]+", osm_value):
            raise ValueError("No OSM, selecione uma categoria ou use uma tag como restaurant ou cafe.")
        query = f'[out:json][timeout:30];nwr["{osm_key}"="{osm_value}"](around:{raio},{lat},{lon});out center;'
        data = consultar("POST", "https://overpass-api.de/api/interpreter", data={"data": query})
        if data.get("remark"):
            raise ValueError("O OSM não concluiu a consulta. Reduza o raio e tente novamente.")
        for p in data.get("elements", []):
            tags = p.get("tags") or {}
            loc = p.get("center") or p
            address = ", ".join(str(tags[k]) for k in ("addr:street", "addr:housenumber", "addr:suburb", "addr:city") if tags.get(k))
            rows.append([tags.get("name", "Sem nome"), address or None, None, None, loc.get("lat"), loc.get("lon")])
    else:
        raise ValueError("Provedor desconhecido.")
    return rows


def detectar_localizacao():
    data = consultar("GET", "https://ipwho.is/")
    if not data.get("success"):
        raise ValueError("Não foi possível detectar sua localização. Informe as coordenadas.")
    lat, lon = coordenadas(data.get("latitude"), data.get("longitude"))
    return lat, lon, data.get("city", "Localização aproximada")


def salvar_excel(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    book = Workbook()
    sheet = book.active
    sheet.title = "Empresas"
    sheet.append(COLUNAS)
    for row in rows:
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        # Dados externos são texto, nunca fórmulas executáveis.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="126B58")
    for col, width in zip("ABCDEF", [36, 65, 14, 14, 18, 18]):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
